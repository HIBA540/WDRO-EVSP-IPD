"""Export Excel des KPIs opérationnels."""
import os
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config.settings import ALL_CASES, CALIBRATION_METHODS


def export_operational_kpis_excel(all_res, all_op_kpis, calib_results, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "Operational KPIs"
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="2C3E50")
    ipd_fill = PatternFill("solid", fgColor="F8E3C9")

    headers = ['Case', 'Method', 'Family', 'ρ',
               'E_OOS (€)', 'CVaR90 (€)', 'Reliability',
               'J_IS (€)', 'ΔJ%', 'Composite Score']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center')

    for cname in ALL_CASES:
        for mk in calib_results.keys():
            res = all_res.get(cname, {}).get(mk, {})
            rho_v = calib_results.get(mk, {}).get('rho', 0.)
            family = CALIBRATION_METHODS.get(mk, {}).get('family', '')
            row = [cname, CALIBRATION_METHODS.get(mk, {}).get('label', mk),
                   family, rho_v,
                   res.get('E_OOS', ''), res.get('CVaR90', ''),
                   res.get('reliability', ''), res.get('J_IS', ''),
                   res.get('delta_J_pct', ''), res.get('composite_score', '')]
            ws.append(row)
            if mk == 'rho_fics':
                for cell in ws[ws.max_row]:
                    cell.fill = ipd_fill
                    cell.font = Font(bold=True)

    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = 16

    path = os.path.join(out_dir, 'operational_kpis_full.xlsx')
    wb.save(path)
    print(f"  [Excel] Saved: {path}")